#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read comprehensive assessment rules document"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from docx import Document

def read_rules():
    doc_path = r"d:\AIOCR\PaddleOCRRAG\data\rules\03、计算机学院综合测评实施细则（2025）.docx"

    doc = Document(doc_path)

    print("=" * 70)
    print("综测实施细则内容")
    print("=" * 70)

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            # 查找A类竞赛相关内容
            if any(keyword in text for keyword in ['A类', '竞赛', '蓝桥杯', '加分', '一等奖', '省级', '国家级']):
                print(f"\n[段落{i}] {text[:300]}")

    # 也检查表格
    print("\n" + "=" * 70)
    print("表格内容")
    print("=" * 70)

    for table_idx, table in enumerate(doc.tables):
        print(f"\n[表格{table_idx}]")
        for row_idx, row in enumerate(table.rows):
            row_text = [cell.text.strip() for cell in row.cells]
            if any('竞赛' in cell or 'A类' in cell or '加分' in cell or '分' in cell for cell in row_text):
                print(f"  行{row_idx}: {' | '.join(row_text)}")

if __name__ == "__main__":
    read_rules()
