"""
综测Excel数据填充服务
负责将RAG检索结果填充到Excel表格
"""
import json
import re
import logging
from typing import Dict, Any, List, Optional
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from app.core.logger import get_logger
from app.services.rag_comprehensive_service import get_rag_comprehensive_service

logger = get_logger(__name__)


class ExcelFillService:
    """Excel数据填充服务"""
    
    COLUMN_MAPPING = {
        'A1—基础分': 'A1',
        'A2—附加分': 'A2',
        'A3—扣分项': 'A3',
        '思想道德素质(A)总分': 'A总分',
        '思想道德素质(A)总分%': 'A加权',
        '学习成绩': 'B原始',
        '学习成绩%': 'B百分比',
        '学习成绩70%': 'B加权',
        'C1—科技竞赛项目': 'C1',
        'C2—体育竞技项目': 'C2',
        'C3—文化类竞赛项目': 'C3',
        'C4—创新创业实践项目': 'C4',
        '素质拓展(C)总分': 'C总分',
        '素质拓展(C)总分10%': 'C加权',
        '综合测评总成绩8%': '总成绩',
        '学生签字': '签字'
    }
    
    HEADER_ROW = 3
    
    def __init__(self):
        """初始化服务"""
        self.rag_service = get_rag_comprehensive_service()
        logger.info("Excel填充服务初始化完成")
    
    def _log_data_trace(self, step: str, data: Any, level: str = "info"):
        """数据追踪日志"""
        log_msg = f"[Excel填充-{step}]"
        
        if isinstance(data, (dict, list)):
            try:
                data_str = json.dumps(data, ensure_ascii=False, indent=2)
                if len(data_str) > 1000:
                    data_str = data_str[:1000] + "..."
            except Exception:
                data_str = str(data)[:1000]
        else:
            data_str = str(data)[:1000]
        
        if level == "error":
            logger.error(f"{log_msg}\n{data_str}")
        elif level == "warning":
            logger.warning(f"{log_msg}\n{data_str}")
        else:
            logger.info(f"{log_msg}\n{data_str}")
    
    async def fill_excel_from_rag(
        self,
        template_path: str,
        output_path: str,
        raw_data: str,
        academic_year: str = None,
        semester: str = None,
        class_id: str = None
    ) -> Dict[str, Any]:
        """使用RAG检索结果填充Excel
        
        Args:
            template_path: 模板文件路径
            output_path: 输出文件路径
            raw_data: 原始数据（OCR识别文本或成绩单内容）
            academic_year: 学年
            semester: 学期
            class_id: 班级ID
            
        Returns:
            填充结果
        """
        self._log_data_trace("开始填充", {
            "template": template_path,
            "output": output_path,
            "academic_year": academic_year,
            "semester": semester,
            "class_id": class_id
        })
        
        try:
            wb = load_workbook(template_path)
            
            if '综合测评计算表' not in wb.sheetnames:
                wb.create_sheet('综合测评计算表')
            if '加减分说明' not in wb.sheetnames:
                wb.create_sheet('加减分说明')
            
            ws_main = wb['综合测评计算表']
            ws_detail = wb['加减分说明']
            
            student_list = self._extract_student_list(ws_main)
            self._log_data_trace("提取学生列表", student_list)
            
            weight_config = await self.rag_service.get_weight_config(
                academic_year or "2024-2025", 
                semester or "1"
            )
            self._log_data_trace("权重配置", weight_config)
            
            fill_result = await self.rag_service.generate_excel_fill_data(
                raw_data, student_list
            )
            
            if not fill_result.get("success"):
                return {
                    "success": False,
                    "error": fill_result.get("error", "生成填充数据失败")
                }
            
            fill_data = fill_result.get("fill_data", [])
            self._log_data_trace("填充数据", fill_data)
            
            fill_stats = {
                "total": len(fill_data),
                "success": 0,
                "failed": 0,
                "errors": []
            }
            
            for student_data in fill_data:
                try:
                    row_index = self._find_student_row(ws_main, student_data.get("student_id"))
                    
                    if row_index:
                        self._fill_student_row(ws_main, row_index, student_data)
                        self._fill_detail_sheet(ws_detail, student_data)
                        fill_stats["success"] += 1
                        
                        self._log_data_trace(f"填充学生成功", {
                            "student_id": student_data.get("student_id"),
                            "row": row_index
                        })
                    else:
                        fill_stats["failed"] += 1
                        fill_stats["errors"].append({
                            "student_id": student_data.get("student_id"),
                            "error": "未找到对应行"
                        })
                        
                except Exception as e:
                    fill_stats["failed"] += 1
                    fill_stats["errors"].append({
                        "student_id": student_data.get("student_id"),
                        "error": str(e)
                    })
                    logger.error(f"填充学生数据失败: {e}")
            
            self._calculate_weighted_scores(ws_main, weight_config)
            
            self._calculate_rankings(ws_main)
            
            wb.save(output_path)
            wb.close()
            
            self._log_data_trace("填充完成", fill_stats)
            
            return {
                "success": True,
                "output_path": output_path,
                "stats": fill_stats,
                "weight_config": weight_config,
                "filled_at": datetime.now().isoformat()
            }
            
        except Exception as e:
            logger.error(f"Excel填充失败: {e}", exc_info=True)
            return {"success": False, "error": str(e)}
    
    def _extract_student_list(self, ws) -> List[Dict]:
        """从工作表提取学生列表"""
        students = []
        
        for row in range(self.HEADER_ROW + 1, ws.max_row + 1):
            student_id = ws.cell(row=row, column=5).value
            student_name = ws.cell(row=row, column=4).value
            
            if student_id and str(student_id).replace('.0', '').isdigit():
                students.append({
                    "row_index": row,
                    "student_id": str(student_id).replace('.0', ''),
                    "student_name": str(student_name) if student_name else ""
                })
        
        return students
    
    def _find_student_row(self, ws, student_id: str) -> Optional[int]:
        """根据学号查找行号"""
        if not student_id:
            return None
        
        student_id = str(student_id).replace('.0', '')
        
        for row in range(self.HEADER_ROW + 1, ws.max_row + 1):
            cell_value = ws.cell(row=row, column=5).value
            if cell_value:
                cell_id = str(cell_value).replace('.0', '')
                if cell_id == student_id:
                    return row
        
        return None
    
    def _fill_student_row(self, ws, row: int, data: Dict):
        """填充学生行数据"""
        columns = data.get("columns", {})
        
        col_map = {
            'A1—基础分': 6,
            'A2—附加分': 7,
            'A3—扣分项': 8,
            '思想道德素质(A)总分': 9,
            '思想道德素质(A)总分%': 10,
            '学习成绩': 11,
            '学习成绩%': 12,
            '学习成绩70%': 13,
            'C1—科技竞赛项目': 14,
            'C2—体育竞技项目': 15,
            'C3—文化类竞赛项目': 16,
            'C4—创新创业实践项目': 17,
            '素质拓展(C)总分': 18,
            '素质拓展(C)总分10%': 19,
            '综合测评总成绩8%': 20,
            '学生签字': 21
        }
        
        for col_name, col_idx in col_map.items():
            value = columns.get(col_name)
            if value is not None:
                try:
                    ws.cell(row=row, column=col_idx).value = float(value)
                except (ValueError, TypeError):
                    ws.cell(row=row, column=col_idx).value = value
    
    def _fill_detail_sheet(self, ws, student_data: Dict):
        """填充加减分明细表
        
        加减分说明表结构：
        - 列1(A): 专业
        - 列2(B): 班级
        - 列3(C): 姓名
        - 列4(D): A1—基础分
        - 列5(E): A2—附加分
        - 列6(F): A3—扣分项
        - 列7(G): C1—科技竞赛项目
        - 列8(H): C2—体育竞技项目
        - 列9(I): C3—文化类竞赛项目
        - 列10(J): C4—创新创业实践项目
        """
        columns = student_data.get("columns", {})
        student_id = student_data.get("student_id", "")
        student_name = student_data.get("student_name", "")
        major = student_data.get("major", "")
        class_name = student_data.get("class_name", "")
        
        row = ws.max_row + 1
        
        ws.cell(row=row, column=1).value = major
        ws.cell(row=row, column=2).value = class_name
        ws.cell(row=row, column=3).value = student_name
        ws.cell(row=row, column=4).value = columns.get('A1—基础分', 0)
        ws.cell(row=row, column=5).value = columns.get('A2—附加分', 0)
        ws.cell(row=row, column=6).value = columns.get('A3—扣分项', 0)
        ws.cell(row=row, column=7).value = columns.get('C1—科技竞赛项目', 0)
        ws.cell(row=row, column=8).value = columns.get('C2—体育竞技项目', 0)
        ws.cell(row=row, column=9).value = columns.get('C3—文化类竞赛项目', 0)
        ws.cell(row=row, column=10).value = columns.get('C4—创新创业实践项目', 0)
        
        logger.debug(f"填充加减分明细: {student_name} - 专业:{major}, 班级:{class_name}")
    
    def _calculate_weighted_scores(self, ws, weight_config: Dict):
        """计算加权分数"""
        config = weight_config.get("config", {})
        a_weight = config.get("a_weight", 20.0) / 100
        b_weight = config.get("b_weight", 70.0) / 100
        c_weight = config.get("c_weight", 10.0) / 100
        
        self._log_data_trace("计算加权分数", {
            "a_weight": a_weight,
            "b_weight": b_weight,
            "c_weight": c_weight
        })
        
        for row in range(self.HEADER_ROW + 1, ws.max_row + 1):
            try:
                a_total = ws.cell(row=row, column=9).value or 0
                b_raw = ws.cell(row=row, column=11).value or 0
                c_total = ws.cell(row=row, column=18).value or 0
                
                a_weighted = float(a_total) * a_weight
                b_percentage = float(b_raw)
                b_weighted = float(b_raw) * b_weight
                c_weighted = float(c_total) * c_weight
                
                ws.cell(row=row, column=10).value = round(a_weighted, 2)
                ws.cell(row=row, column=12).value = round(b_percentage, 2)
                ws.cell(row=row, column=13).value = round(b_weighted, 2)
                ws.cell(row=row, column=19).value = round(c_weighted, 2)
                
                total_score = a_weighted + b_weighted + c_weighted
                ws.cell(row=row, column=20).value = round(total_score, 2)
                
            except Exception as e:
                logger.warning(f"计算第{row}行加权分数失败: {e}")
    
    def _calculate_rankings(self, ws):
        """计算排名"""
        scores = []
        
        for row in range(self.HEADER_ROW + 1, ws.max_row + 1):
            total_score = ws.cell(row=row, column=20).value
            if total_score is not None:
                try:
                    scores.append((row, float(total_score)))
                except (ValueError, TypeError):
                    pass
        
        scores.sort(key=lambda x: x[1], reverse=True)
        
        for rank, (row, score) in enumerate(scores, 1):
            ws.cell(row=row, column=1).value = rank
    
    async def process_ocr_result(
        self,
        ocr_text: str,
        student_info: Dict = None,
        academic_year: str = None,
        semester: str = None
    ) -> Dict[str, Any]:
        """处理OCR识别结果并返回填充数据
        
        Args:
            ocr_text: OCR识别文本
            student_info: 学生信息
            academic_year: 学年
            semester: 学期
            
        Returns:
            处理结果
        """
        self._log_data_trace("处理OCR结果-开始", {
            "ocr_text_length": len(ocr_text),
            "student_info": student_info
        })
        
        try:
            result = await self.rag_service.analyze_certificate(
                ocr_text, student_info
            )
            
            if result.get("success"):
                self._log_data_trace("OCR处理成功", result)
                
                return {
                    "success": True,
                    "category": result.get("category"),
                    "sub_category": result.get("sub_category"),
                    "score": result.get("score"),
                    "level": result.get("level"),
                    "certificate_name": result.get("certificate_name"),
                    "explanation": result.get("explanation"),
                    "rules_matched": result.get("rules_matched", [])
                }
            
            return result
            
        except Exception as e:
            logger.error(f"处理OCR结果失败: {e}", exc_info=True)
            return {"success": False, "error": str(e)}
    
    async def batch_process_students(
        self,
        students_data: List[Dict],
        academic_year: str,
        semester: str
    ) -> Dict[str, Any]:
        """批量处理学生数据
        
        Args:
            students_data: 学生数据列表
            academic_year: 学年
            semester: 学期
            
        Returns:
            处理结果
        """
        self._log_data_trace("批量处理学生-开始", {
            "student_count": len(students_data),
            "academic_year": academic_year,
            "semester": semester
        })
        
        results = []
        success_count = 0
        failed_count = 0
        
        weight_config = await self.rag_service.get_weight_config(academic_year, semester)
        
        for student in students_data:
            try:
                result = await self.rag_service.calculate_student_score(
                    student_id=student.get("student_id"),
                    student_name=student.get("student_name"),
                    class_name=student.get("class_name", ""),
                    academic_info=student.get("academic_info", {}),
                    certificate_info=student.get("certificate_info", []),
                    score_details=student.get("score_details", [])
                )
                
                if result.get("success"):
                    success_count += 1
                else:
                    failed_count += 1
                
                results.append(result)
                
            except Exception as e:
                failed_count += 1
                results.append({
                    "success": False,
                    "student_id": student.get("student_id"),
                    "error": str(e)
                })
        
        self._log_data_trace("批量处理完成", {
            "success": success_count,
            "failed": failed_count
        })
        
        return {
            "success": True,
            "total": len(students_data),
            "success_count": success_count,
            "failed_count": failed_count,
            "results": results,
            "weight_config": weight_config
        }


excel_fill_service = ExcelFillService()


def get_excel_fill_service() -> ExcelFillService:
    """获取Excel填充服务实例"""
    return excel_fill_service
