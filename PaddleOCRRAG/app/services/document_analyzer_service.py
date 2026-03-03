"""
文档分析服务模块
提供文档读取、AI检索提取规则、比例数据结构化提取等功能
"""
import os
import re
import time
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from app.core.logger import get_logger, rag_logger, track_performance
from app.core.llm_manager import LLMManager

logger = get_logger(__name__)


class DocumentAnalyzerService:
    """文档分析服务类"""
    
    DEFAULT_DOC_PATH = r"D:\PaddleOCR\PaddleOCRRAG\data\rules\03、计算机学院综合测评实施细则（2025）.docx"
    
    def __init__(self):
        """初始化文档分析服务"""
        self._llm_manager = None
        self._vector_db = None
        self._document_cache = {}
        logger.info("[DocumentAnalyzerService] 文档分析服务初始化完成")
    
    @property
    def llm_manager(self):
        """延迟加载LLM管理器"""
        if self._llm_manager is None:
            self._llm_manager = LLMManager()
            logger.debug("[DocumentAnalyzerService] LLM管理器已加载")
        return self._llm_manager
    
    @property
    def vector_db(self):
        """延迟加载向量数据库"""
        if self._vector_db is None:
            from app.rag.vector_db.vector_db import get_vector_db
            self._vector_db = get_vector_db()
            logger.debug("[DocumentAnalyzerService] 向量数据库已加载")
        return self._vector_db
    
    @track_performance("read_document")
    def read_document(self, file_path: str = None) -> Dict[str, Any]:
        """
        读取指定docx文档
        
        Args:
            file_path: 文档路径，默认使用综测规则文档
            
        Returns:
            包含文档内容的字典
        """
        start_time = time.time()
        doc_path = file_path or self.DEFAULT_DOC_PATH
        
        logger.info(f"[read_document] 开始读取文档: {doc_path}")
        
        try:
            if not os.path.exists(doc_path):
                error_msg = f"文档不存在: {doc_path}"
                logger.error(f"[read_document] {error_msg}")
                return {
                    "success": False,
                    "error": error_msg,
                    "content": "",
                    "tables": [],
                    "paragraphs": []
                }
            
            from docx import Document
            doc = Document(doc_path)
            
            paragraphs = []
            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if text:
                    paragraphs.append({
                        "index": i,
                        "text": text,
                        "style": para.style.name if para.style else None
                    })
            
            tables = []
            for i, table in enumerate(doc.tables):
                table_data = self._extract_table_data(table, i)
                if table_data:
                    tables.append(table_data)
            
            full_text = "\n".join([p["text"] for p in paragraphs])
            
            duration_ms = int((time.time() - start_time) * 1000)
            logger.info(
                f"[read_document] 文档读取完成: 段落数={len(paragraphs)}, 表格数={len(tables)}",
                extra={
                    'duration_ms': duration_ms,
                    'params': {
                        'file_path': doc_path,
                        'paragraph_count': len(paragraphs),
                        'table_count': len(tables),
                        'total_chars': len(full_text)
                    }
                }
            )
            
            return {
                "success": True,
                "error": None,
                "content": full_text,
                "tables": tables,
                "paragraphs": paragraphs,
                "metadata": {
                    "file_path": doc_path,
                    "file_name": os.path.basename(doc_path),
                    "paragraph_count": len(paragraphs),
                    "table_count": len(tables),
                    "total_chars": len(full_text)
                }
            }
            
        except ImportError as e:
            error_msg = f"缺少依赖库: {str(e)}"
            logger.error(f"[read_document] {error_msg}")
            return {
                "success": False,
                "error": error_msg,
                "content": "",
                "tables": [],
                "paragraphs": []
            }
        except Exception as e:
            error_msg = f"读取文档失败: {str(e)}"
            logger.error(f"[read_document] {error_msg}", exc_info=True)
            return {
                "success": False,
                "error": error_msg,
                "content": "",
                "tables": [],
                "paragraphs": []
            }
    
    def _extract_table_data(self, table, table_index: int) -> Optional[Dict[str, Any]]:
        """
        提取表格数据
        
        Args:
            table: docx表格对象
            table_index: 表格索引
            
        Returns:
            表格数据字典
        """
        try:
            if not table.rows:
                return None
            
            rows_data = []
            for row in table.rows:
                cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                rows_data.append(cells)
            
            if not rows_data:
                return None
            
            headers = rows_data[0] if rows_data else []
            
            logger.debug(f"[_extract_table_data] 表格{table_index}: {len(rows_data)}行, 列头={headers[:3]}...")
            
            return {
                "index": table_index,
                "headers": headers,
                "rows": rows_data,
                "row_count": len(rows_data),
                "col_count": len(headers) if headers else 0
            }
            
        except Exception as e:
            logger.warning(f"[_extract_table_data] 提取表格{table_index}失败: {e}")
            return None
    
    @track_performance("retrieve_rules_by_rag")
    def retrieve_rules_by_rag(self, query: str = None, top_k: int = 10) -> Dict[str, Any]:
        """
        使用RAG系统检索综测计算规则
        
        Args:
            query: 检索查询，默认检索综测计算规则
            top_k: 返回结果数量
            
        Returns:
            检索结果字典
        """
        start_time = time.time()
        
        if query is None:
            query = "综测计算规则 综合测评分数计算方法 品德行为分 学习成绩分 素质拓展分"
        
        logger.info(f"[retrieve_rules_by_rag] 开始RAG检索: query='{query[:50]}...', top_k={top_k}")
        
        try:
            results = self.vector_db.search_relevant(
                query=query,
                top_k=top_k,
                similarity_threshold=0.55
            )
            
            documents = results.get("documents", [])
            metadatas = results.get("metadatas", [])
            distances = results.get("distances", [])
            
            rules = []
            for i, (doc, meta, dist) in enumerate(zip(documents, metadatas, distances)):
                similarity = 1 / (1 + dist)
                rule_item = {
                    "index": i,
                    "content": doc,
                    "metadata": meta,
                    "distance": dist,
                    "similarity": round(similarity, 4)
                }
                rules.append(rule_item)
                logger.debug(f"[retrieve_rules_by_rag] 规则{i}: 相似度={similarity:.4f}, 内容前50字='{doc[:50]}...'")
            
            duration_ms = int((time.time() - start_time) * 1000)
            logger.info(
                f"[retrieve_rules_by_rag] RAG检索完成: 返回{len(rules)}条规则",
                extra={
                    'duration_ms': duration_ms,
                    'params': {
                        'query': query[:100] if len(query) > 100 else query,
                        'top_k': top_k,
                        'result_count': len(rules)
                    }
                }
            )
            
            rag_logger.log_retrieval(query, {}, len(rules), duration_ms)
            
            return {
                "success": True,
                "rules": rules,
                "total": len(rules),
                "query": query
            }
            
        except Exception as e:
            error_msg = f"RAG检索失败: {str(e)}"
            logger.error(f"[retrieve_rules_by_rag] {error_msg}", exc_info=True)
            return {
                "success": False,
                "rules": [],
                "total": 0,
                "error": error_msg,
                "query": query
            }
    
    @track_performance("extract_ratio_data")
    def extract_ratio_data(self, content: str = None, use_rag: bool = True) -> Dict[str, Any]:
        """
        提取综测计算比例数据
        
        提取公式：M=20%*A+70%*B+10%*C
        各部分含义：A—品德行为分，B—学习成绩分，C—素质拓展分
        
        Args:
            content: 文档内容，如果为None则从默认文档读取
            use_rag: 是否使用RAG辅助提取
            
        Returns:
            比例数据字典
        """
        start_time = time.time()
        logger.info("[extract_ratio_data] 开始提取综测计算比例数据")
        
        try:
            if content is None:
                doc_result = self.read_document()
                if not doc_result.get("success"):
                    return {
                        "success": False,
                        "error": doc_result.get("error", "无法读取文档"),
                        "formula": None,
                        "ratios": {},
                        "components": {}
                    }
                content = doc_result.get("content", "")
            
            formula_pattern = r'M\s*=\s*(\d+)%?\s*\*?\s*A\s*\+\s*(\d+)%?\s*\*?\s*B\s*\+\s*(\d+)%?\s*\*?\s*C'
            formula_match = re.search(formula_pattern, content, re.IGNORECASE)
            
            ratios = {}
            formula = None
            components = {}
            
            if formula_match:
                ratio_a = int(formula_match.group(1))
                ratio_b = int(formula_match.group(2))
                ratio_c = int(formula_match.group(3))
                
                formula = f"M={ratio_a}%*A+{ratio_b}%*B+{ratio_c}%*C"
                
                ratios = {
                    "A": ratio_a / 100,
                    "B": ratio_b / 100,
                    "C": ratio_c / 100
                }
                
                logger.debug(f"[extract_ratio_data] 找到公式: {formula}")
            else:
                logger.debug("[extract_ratio_data] 未找到标准公式格式，尝试其他模式")
                alt_pattern = r'综合测评.*?(\d+)%.*?(\d+)%.*?(\d+)%'
                alt_match = re.search(alt_pattern, content, re.IGNORECASE | re.DOTALL)
                if alt_match:
                    ratio_a = int(alt_match.group(1))
                    ratio_b = int(alt_match.group(2))
                    ratio_c = int(alt_match.group(3))
                    formula = f"M={ratio_a}%*A+{ratio_b}%*B+{ratio_c}%*C"
                    ratios = {
                        "A": ratio_a / 100,
                        "B": ratio_b / 100,
                        "C": ratio_c / 100
                    }
                    logger.debug(f"[extract_ratio_data] 通过备选模式找到比例: {ratios}")
            
            component_patterns = [
                (r'A[—\-:：]\s*品德行为分', 'A', '品德行为分'),
                (r'B[—\-:：]\s*学习成绩分', 'B', '学习成绩分'),
                (r'C[—\-:：]\s*素质拓展分', 'C', '素质拓展分'),
                (r'品德行为分.*?占.*?(\d+)%', 'A', '品德行为分'),
                (r'学习成绩分.*?占.*?(\d+)%', 'B', '学习成绩分'),
                (r'素质拓展分.*?占.*?(\d+)%', 'C', '素质拓展分'),
            ]
            
            for pattern, key, name in component_patterns:
                match = re.search(pattern, content, re.IGNORECASE)
                if match and key not in components:
                    components[key] = {
                        "name": name,
                        "description": f"{name}，占比{int(ratios.get(key, 0) * 100)}%"
                    }
                    logger.debug(f"[extract_ratio_data] 识别到组件: {key}={name}")
            
            if use_rag and not ratios:
                logger.info("[extract_ratio_data] 本地提取失败，尝试RAG辅助提取")
                rag_result = self.retrieve_rules_by_rag("综测计算比例 品德行为分 学习成绩分 素质拓展分 占比", top_k=5)
                if rag_result.get("success") and rag_result.get("rules"):
                    for rule in rag_result["rules"]:
                        rule_content = rule.get("content", "")
                        if not ratios:
                            match = re.search(formula_pattern, rule_content, re.IGNORECASE)
                            if match:
                                ratio_a = int(match.group(1))
                                ratio_b = int(match.group(2))
                                ratio_c = int(match.group(3))
                                formula = f"M={ratio_a}%*A+{ratio_b}%*B+{ratio_c}%*C"
                                ratios = {
                                    "A": ratio_a / 100,
                                    "B": ratio_b / 100,
                                    "C": ratio_c / 100
                                }
                                logger.debug(f"[extract_ratio_data] RAG辅助找到公式: {formula}")
                                break
            
            if not components:
                components = {
                    "A": {"name": "品德行为分", "description": "品德行为分，占比20%"},
                    "B": {"name": "学习成绩分", "description": "学习成绩分，占比70%"},
                    "C": {"name": "素质拓展分", "description": "素质拓展分，占比10%"}
                }
            
            duration_ms = int((time.time() - start_time) * 1000)
            logger.info(
                f"[extract_ratio_data] 比例数据提取完成: formula={formula}",
                extra={
                    'duration_ms': duration_ms,
                    'params': {
                        'formula': formula,
                        'ratios': ratios,
                        'components_count': len(components)
                    }
                }
            )
            
            return {
                "success": True,
                "formula": formula,
                "ratios": ratios,
                "components": components,
                "raw_content_length": len(content) if content else 0
            }
            
        except Exception as e:
            error_msg = f"提取比例数据失败: {str(e)}"
            logger.error(f"[extract_ratio_data] {error_msg}", exc_info=True)
            return {
                "success": False,
                "error": error_msg,
                "formula": None,
                "ratios": {},
                "components": {}
            }
    
    @track_performance("generate_analysis_description")
    def generate_analysis_description(self, content: str = None, max_length: int = 100) -> str:
        """
        生成文档分析说明
        
        Args:
            content: 文档内容，如果为None则从默认文档读取
            max_length: 最大长度限制（默认100字）
            
        Returns:
            文档分析说明文本
        """
        start_time = time.time()
        logger.info(f"[generate_analysis_description] 开始生成文档分析说明，最大长度={max_length}字")
        
        try:
            if content is None:
                doc_result = self.read_document()
                if not doc_result.get("success"):
                    return "文档分析失败：无法读取文档内容"
                content = doc_result.get("content", "")
            
            title = ""
            title_patterns = [
                r'([^。\n]{5,30}实施细则)',
                r'([^。\n]{5,30}综合测评)',
                r'([^。\n]{5,30}办法)',
            ]
            for pattern in title_patterns:
                match = re.search(pattern, content)
                if match:
                    title = match.group(1).strip()
                    break
            
            if not title:
                title = "计算机学院综合测评实施细则"
            
            keywords = ["综合测评", "品德行为", "学习成绩", "素质拓展", "加分", "扣分", "竞赛", "证书"]
            found_keywords = [kw for kw in keywords if kw in content]
            
            analysis = f"《{title}》规定了学生综合测评的计算方法，"
            
            if found_keywords:
                analysis += f"涵盖{'+'.join(found_keywords[:4])}等内容，"
            
            analysis += "用于指导学生综合素质评价工作。"
            
            if len(analysis) > max_length:
                analysis = analysis[:max_length-3] + "..."
            
            duration_ms = int((time.time() - start_time) * 1000)
            logger.info(
                f"[generate_analysis_description] 分析说明生成完成: 长度={len(analysis)}字",
                extra={
                    'duration_ms': duration_ms,
                    'params': {
                        'analysis_length': len(analysis),
                        'max_length': max_length,
                        'title': title
                    }
                }
            )
            
            return analysis
            
        except Exception as e:
            error_msg = f"生成分析说明失败: {str(e)}"
            logger.error(f"[generate_analysis_description] {error_msg}", exc_info=True)
            return f"文档分析说明生成失败: {str(e)}"
    
    @track_performance("analyze_document")
    def analyze_document(self, file_path: str = None) -> Dict[str, Any]:
        """
        完整分析文档，返回结构化数据
        
        返回格式：{rules: [...], ratios: {...}, analysis: "..."}
        
        Args:
            file_path: 文档路径，默认使用综测规则文档
            
        Returns:
            结构化分析结果
        """
        start_time = time.time()
        doc_path = file_path or self.DEFAULT_DOC_PATH
        
        logger.info(f"[analyze_document] 开始完整文档分析: {doc_path}")
        
        try:
            doc_result = self.read_document(doc_path)
            if not doc_result.get("success"):
                return {
                    "success": False,
                    "error": doc_result.get("error", "文档读取失败"),
                    "rules": [],
                    "ratios": {},
                    "analysis": ""
                }
            
            content = doc_result.get("content", "")
            
            rules_result = self.retrieve_rules_by_rag()
            rules = rules_result.get("rules", [])
            
            ratios_result = self.extract_ratio_data(content)
            ratios = {
                "formula": ratios_result.get("formula"),
                "data": ratios_result.get("ratios", {}),
                "components": ratios_result.get("components", {})
            }
            
            analysis = self.generate_analysis_description(content)
            
            result = {
                "success": True,
                "rules": rules,
                "ratios": ratios,
                "analysis": analysis,
                "metadata": {
                    "file_path": doc_path,
                    "file_name": os.path.basename(doc_path),
                    "rules_count": len(rules),
                    "has_formula": ratios.get("formula") is not None,
                    "analysis_length": len(analysis)
                }
            }
            
            duration_ms = int((time.time() - start_time) * 1000)
            logger.info(
                f"[analyze_document] 文档分析完成: rules={len(rules)}, has_formula={ratios.get('formula') is not None}",
                extra={
                    'duration_ms': duration_ms,
                    'params': {
                        'file_path': doc_path,
                        'rules_count': len(rules),
                        'has_formula': ratios.get("formula") is not None,
                        'analysis_length': len(analysis)
                    }
                }
            )
            
            return result
            
        except Exception as e:
            error_msg = f"文档分析失败: {str(e)}"
            logger.error(f"[analyze_document] {error_msg}", exc_info=True)
            return {
                "success": False,
                "error": error_msg,
                "rules": [],
                "ratios": {},
                "analysis": ""
            }


_document_analyzer_service = None


def get_document_analyzer_service() -> DocumentAnalyzerService:
    """
    获取文档分析服务单例实例
    
    Returns:
        DocumentAnalyzerService实例
    """
    global _document_analyzer_service
    if _document_analyzer_service is None:
        _document_analyzer_service = DocumentAnalyzerService()
        logger.debug("[get_document_analyzer_service] 创建文档分析服务单例")
    return _document_analyzer_service
