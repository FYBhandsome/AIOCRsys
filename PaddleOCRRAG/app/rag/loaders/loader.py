import os
from pathlib import Path
from app.core.config_manager import settings


class RuleDocumentLoader:
    """综测规则文档加载器"""
    
    def __init__(self):
        self.documents = []
        self.chunk_size = 1000
        self.chunk_overlap = 150
    
    def load_all_documents(self):
        """加载所有规则文档"""
        if not os.path.exists(settings.RULES_DOCS_PATH):
            os.makedirs(settings.RULES_DOCS_PATH, exist_ok=True)
            print(f"规则文档目录不存在，已创建: {settings.RULES_DOCS_PATH}")
            return self
        
        loaded_count = 0
        for filename in os.listdir(settings.RULES_DOCS_PATH):
            file_path = os.path.join(settings.RULES_DOCS_PATH, filename)
            
            if filename.startswith('.'):
                continue
            
            try:
                if filename.endswith(".txt"):
                    docs = self._load_txt(file_path)
                elif filename.endswith(".docx"):
                    docs = self._load_docx(file_path)
                elif filename.endswith(".pdf"):
                    docs = self._load_pdf(file_path)
                else:
                    continue
                
                if docs:
                    self.documents.extend(docs)
                    loaded_count += 1
                    print(f"已加载文档: {filename} ({len(docs)} 个片段)")
            except Exception as e:
                print(f"加载文档 {filename} 时出错: {str(e)}")
                continue
        
        print(f"共加载 {loaded_count} 个文档")
        return self
    
    def _load_txt(self, file_path: str):
        """加载文本文件"""
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        return [{'page_content': text, 'metadata': {'source': file_path}}]
    
    def _load_docx(self, file_path: str):
        """加载Word文档"""
        try:
            from docx import Document
            doc = Document(file_path)
            
            text_parts = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    text_parts.append(text)
            
            for table in doc.tables:
                table_text = self._extract_table(table)
                if table_text:
                    text_parts.append(table_text)
            
            full_text = "\n\n".join(text_parts)
            return [{'page_content': full_text, 'metadata': {'source': file_path}}]
            
        except ImportError:
            import docx2txt
            text = docx2txt.process(file_path)
            return [{'page_content': text, 'metadata': {'source': file_path}}]
    
    def _load_pdf(self, file_path: str):
        """加载PDF文件"""
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
            full_text = "\n\n".join(text_parts)
            return [{'page_content': full_text, 'metadata': {'source': file_path}}]
        except ImportError:
            print("pdfplumber未安装，无法加载PDF文件")
            return []
    
    def _extract_table(self, table) -> str:
        """提取表格内容"""
        if not table.rows:
            return ""
        
        rows_data = []
        for row in table.rows:
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            rows_data.append(cells)
        
        if not rows_data:
            return ""
        
        headers = rows_data[0] if rows_data else []
        result_lines = []
        
        if len(rows_data) > 1 and len(headers) >= 2:
            for row in rows_data[1:]:
                if len(row) >= 2:
                    row_text = " | ".join([f"{h}: {v}" for h, v in zip(headers, row) if h and v])
                    if row_text:
                        result_lines.append(row_text)
        else:
            for row in rows_data:
                row_text = " | ".join([v for v in row if v])
                if row_text:
                    result_lines.append(row_text)
        
        return "\n".join(result_lines)
    
    def split_documents(self):
        """分割文档为小块"""
        if not self.documents:
            return []
        
        chunks = []
        for doc in self.documents:
            text = doc.get('page_content', '')
            if not text:
                continue
            
            doc_chunks = self._split_text(text, doc.get('metadata', {}))
            chunks.extend(doc_chunks)
        
        return chunks
    
    def _split_text(self, text: str, metadata: dict) -> list:
        """分割文本为小块"""
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + self.chunk_size
            chunk_text = text[start:end]
            
            if end < len(text):
                for sep in ['。', '！', '？', '\n', '，', ' ']:
                    last_sep = chunk_text.rfind(sep)
                    if last_sep > self.chunk_size // 2:
                        end = start + last_sep + 1
                        chunk_text = text[start:end]
                        break
            
            chunk = {
                'page_content': chunk_text.strip(),
                'metadata': {**metadata, 'chunk_index': len(chunks)}
            }
            chunks.append(chunk)
            
            start = end - self.chunk_overlap
            if start < 0:
                start = 0
        
        return chunks
